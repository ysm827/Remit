"""Paper post-processing plus reproducible LaTeX/PDF delivery."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
from uuid import uuid4

from app.services.async_io import check_cancelled, WorkCancelled
from app.utils.tex_process import run_xelatex
from app.utils.pandoc_process import convert_text
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.config.setting import settings
from app.schemas.enums import CompTemplate
from app.utils.log_util import logger

IMAGE_BLOCK_RE = re.compile(r"^\s*!\[(?P<alt>.*?)\]\((?P<src>.*?)\)\s*$")
HEADING_RE = re.compile(r"^\s*#{1,6}\s+")
ABSTRACT_RE = re.compile(r"^\s*#{1,6}\s*摘要\s*$")
KEYWORD_RE = re.compile(r"^\s*\*{0,2}\s*关(?:键)?词[:：]?\s*\*{0,2}")
QUESTION_LEAD_RE = re.compile(
    r"^(?P<prefix>针对)?(?P<lead>问题[一二三四五六七八九十])[,，：: ]*(?P<body>.*)$",
)

# ---- 中文赛事论文组装（摘要页 / 标题层级 / 参考文献著录） ----

CN_NUM_CHARS = "一二三四五六七八九十"
CHINA_ABSTRACT_RE = re.compile(r"^\s{0,3}#{1,6}\s*摘\s*要\s*$")
REFS_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s*参考文献\s*$")
REF_DEF_RE = re.compile(r"^\s{0,3}\[\^(\d+)\]:\s*(.*)$")
FOOTNOTE_REF_RE = re.compile(r"\[\^(\d+)\]")

# 写作手常直接输入 Unicode 希腊字母；正文字体不含这些字形时 XeLaTeX 会静默丢字，
# 统一换成数学命令交给数学字体渲染。
GREEK_UNICODE_MAP = {
    "α": r"\alpha",
    "β": r"\beta",
    "γ": r"\gamma",
    "δ": r"\delta",
    "ε": r"\varepsilon",
    "ζ": r"\zeta",
    "η": r"\eta",
    "θ": r"\theta",
    "ι": r"\iota",
    "κ": r"\kappa",
    "λ": r"\lambda",
    "μ": r"\mu",
    "ν": r"\nu",
    "ξ": r"\xi",
    "π": r"\pi",
    "ρ": r"\rho",
    "σ": r"\sigma",
    "τ": r"\tau",
    "υ": r"\upsilon",
    "φ": r"\varphi",
    "χ": r"\chi",
    "ψ": r"\psi",
    "ω": r"\omega",
    "Γ": r"\Gamma",
    "Δ": r"\Delta",
    "Θ": r"\Theta",
    "Λ": r"\Lambda",
    "Ξ": r"\Xi",
    "Π": r"\Pi",
    "Σ": r"\Sigma",
    "Φ": r"\Phi",
    "Ψ": r"\Psi",
    "Ω": r"\Omega",
    "ϵ": r"\epsilon",
    "ϑ": r"\vartheta",
    "ϕ": r"\phi",
    "ς": r"\varsigma",
    "ϱ": r"\varrho",
    "ϖ": r"\varpi",
    # 微符号与上/下标：宋体没有这些字形，文本模式会静默丢字（豆腐块）。
    "µ": r"\mu",
    "−": r"-",
    "⁰": r"{}^{0}",
    "¹": r"{}^{1}",
    "²": r"{}^{2}",
    "³": r"{}^{3}",
    "⁴": r"{}^{4}",
    "⁵": r"{}^{5}",
    "⁶": r"{}^{6}",
    "⁷": r"{}^{7}",
    "⁸": r"{}^{8}",
    "⁹": r"{}^{9}",
    "⁺": r"{}^{+}",
    "⁻": r"{}^{-}",
    "₀": r"{}_{0}",
    "₁": r"{}_{1}",
    "₂": r"{}_{2}",
    "₃": r"{}_{3}",
    "₄": r"{}_{4}",
    "₅": r"{}_{5}",
    "₆": r"{}_{6}",
    "₇": r"{}_{7}",
    "₈": r"{}_{8}",
    "₉": r"{}_{9}",
    "₊": r"{}_{+}",
    "₋": r"{}_{-}",
    "≤": r"\leq",
    "≥": r"\geq",
    "≠": r"\neq",
    "≈": r"\approx",
    "∈": r"\in",
    "∞": r"\infty",
}
GREEK_CHAR_RE = re.compile("[" + "".join(GREEK_UNICODE_MAP) + "]")
MATH_SPAN_RE = re.compile(r"(\$\$.+?\$\$|\$[^$\n]+\$|\\\(.+?\\\)|\\\[.+?\\\])", re.S)
MD_HEADING_LINE_RE = re.compile(r"^(\s{0,3}#{1,6})\s+(.*)$")
VERBATIM_BLOCK_RE = re.compile(r"\\begin\{verbatim\}.*?\\end\{verbatim\}", re.DOTALL)
DISPLAY_MATH_RE = re.compile(r"\\\[(.*?)\\\]", re.DOTALL)
REFS_PLACEHOLDER = "REMITREFERENCESECTIONPLACEHOLDER"

_LEGACY_PAPER_OUTPUTS = (
    "res.md",
    "res_polished.md",
    "res.docx",
    "res_polished.docx",
    "res_polished.pdf",
    "paper_render.html",
    "paper_pdf_header.tex",
    "paper_reference.docx",
)


class PaperRenderError(RuntimeError):
    """LaTeX 生成、编译或 PDF 复核失败。"""


@dataclass(frozen=True)
class PaperDeliverables:
    """由同一份可编译 LaTeX 源码生成的最终交付物。"""

    tex_path: Path
    pdf_path: Path
    report_path: Path
    page_count: int


def render_paper_deliverables(
    markdown: str,
    work_dir: str | Path,
    comp_template: CompTemplate,
) -> PaperDeliverables:
    """将已校验终稿转换为 LaTeX，并由该源码编译、复核 PDF。"""
    root = Path(work_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    build_dir = root / ".remit" / "latex-build" / uuid4().hex
    build_dir.mkdir(parents=True, exist_ok=True)
    tex_path = root / "res.tex"
    pdf_path = root / "res.pdf"
    report_path = root / "paper_delivery_report.json"
    candidate_tex = build_dir / "res.tex"
    check_cancelled()
    _convert_markdown_to_latex(markdown, candidate_tex, root, build_dir, comp_template)
    check_cancelled()
    _validate_latex_source(candidate_tex)
    engine = _compile_latex(candidate_tex, build_dir, resource_path=root)
    built_pdf = build_dir / "res.pdf"
    if not built_pdf.is_file():
        raise PaperRenderError("LaTeX 编译命令成功退出，但没有生成 res.pdf")
    pdf_metrics = inspect_pdf_artifact(
        built_pdf,
        comp_template=comp_template,
        minimum_pages=settings.PAPER_MIN_PDF_PAGES,
    )
    report = {
        "status": "pass",
        "source": tex_path.name,
        "pdf": pdf_path.name,
        "compiler": engine,
        "compile_passes": 2,
        "tex_sha256": _sha256(candidate_tex),
        "pdf_sha256": _sha256(built_pdf),
        **pdf_metrics,
    }
    candidate_report = build_dir / report_path.name
    candidate_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    check_cancelled()
    # Compile and inspect completely before replacing any previous deliverable.
    # Keep rollback copies in this build if a final replacement fails (e.g. PDF locked).
    replacements = [
        (candidate_tex, tex_path),
        (built_pdf, pdf_path),
        (candidate_report, report_path),
    ]
    originals = {}
    for _, target in replacements:
        if target.exists():
            backup = build_dir / (target.name + ".previous")
            shutil.copy2(target, backup)
            originals[target] = backup
    published = []
    try:
        for candidate, target in replacements:
            temporary = target.with_name(f".{target.name}.{build_dir.name}.tmp")
            shutil.copy2(candidate, temporary)
            temporary.replace(target)
            published.append(target)
    except OSError:
        for target in reversed(published):
            if target in originals:
                shutil.copy2(originals[target], target)
            else:
                target.unlink(missing_ok=True)
        raise
    _remove_legacy_paper_outputs(root)
    return PaperDeliverables(
        tex_path=tex_path,
        pdf_path=pdf_path,
        report_path=report_path,
        page_count=int(pdf_metrics["page_count"]),
    )


def _convert_markdown_to_latex(
    markdown: str,
    output_path: Path,
    resource_path: Path,
    build_dir: Path,
    comp_template: CompTemplate,
    *,
    mode: str = "full_paper",
) -> None:
    if comp_template == CompTemplate.CHINA:
        _assemble_china_paper(markdown, output_path, resource_path, mode=mode)
        return

    markdown = _extract_inline_images(markdown)
    header_path = build_dir / "paper_header.tex"
    header_path.write_text(
        build_pdf_header(resource_path, comp_template), encoding="utf-8"
    )
    try:
        convert_text(
            markdown,
            to="latex",
            format="markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_html",
            outputfile=str(output_path),
            extra_args=[
                f"--resource-path={resource_path}",
                f"--include-in-header={header_path}",
                "--standalone",
                "--wrap=none",
            ],
        )
    except WorkCancelled:
        raise
    except Exception as exc:
        raise PaperRenderError(f"Pandoc 生成 LaTeX 失败: {exc}") from exc


def _escape_latex_text(text: str) -> str:
    """转义纯文本片段，让标题、关键词与参考文献条目安全进入 LaTeX。"""
    text = text.replace("**", "")
    out: list[str] = []
    for char in text:
        if char == "\\":
            out.append(r"\textbackslash{}")
        elif char in "&%$#_{}":
            out.append("\\" + char)
        elif char == "~":
            out.append(r"\textasciitilde{}")
        elif char == "^":
            out.append(r"\textasciicircum{}")
        else:
            out.append(char)
    return "".join(out)


def _normalize_keywords(keywords: str) -> str:
    """关键词统一用顿号分隔；写作手可能用空格、逗号或分号分隔。"""
    parts = [
        part.strip()
        for part in re.split(r"[、，,；;]|\s{2,}|\t", keywords)
        if part.strip()
    ]
    return "、".join(parts) if parts else keywords.strip()


def _normalize_heading_level(marks: str, title: str) -> str:
    """按作者手写编号推断标题真实层级，修正跨章节 # 数量不一致。"""
    text = title.strip()
    if re.match(rf"^[{CN_NUM_CHARS}]{{1,3}}、", text):
        return "#"
    match = re.match(r"^\d+(?:\.\d+){1,2}(?=\D|$)", text)
    if match:
        return "#" * min(match.group(0).count(".") + 1, 6)
    return marks.strip()


def _split_china_paper(
    markdown: str,
) -> tuple[str, str, str, list[tuple[int, str]], str]:
    """把整篇 Markdown 拆成标题、摘要、关键词、参考文献条目与正文。

    Writer 各章节带手工编号（一、/1.1/5.1.1），引用为 [^n] 脚注语法；
    这里只拆分结构，不改动任何句子内容。索引先定位后一次性切片，
    避免边删边改导致的下标漂移。
    """
    markdown = re.sub(
        r"(?m)^[ \t]*(?:\*\*)?摘[ \t]*要(?:\*\*)?[ \t]*$", "## 摘要", markdown
    )
    lines = markdown.splitlines()
    headings = [
        (idx, line)
        for idx, line in enumerate(lines)
        if re.match(r"^\s{0,3}#{1,6}\s+", line)
    ]
    abstract_idx = next(
        (i for i, line in headings if CHINA_ABSTRACT_RE.match(line)), None
    )
    refs_idx = next((i for i, line in headings if REFS_HEADING_RE.match(line)), None)

    title, title_idx = "", None
    limit = abstract_idx if abstract_idx is not None else len(lines)
    for idx, line in enumerate(lines[:limit]):
        match = re.match(r"^\s{0,3}#\s+(.+?)\s*$", line)
        if match:
            title, title_idx = match.group(1).strip(), idx
            break
    if title_idx is None:
        for idx, line in headings:
            if idx in (abstract_idx, refs_idx):
                continue
            match = re.match(r"^\s{0,3}#\s+(.+?)\s*$", line)
            if match:
                title, title_idx = match.group(1).strip(), idx
                break

    abstract_lines: list[str] = []
    keywords = ""
    abstract_end: int | None = None
    if abstract_idx is not None:
        abstract_end = next((i for i, _ in headings if i > abstract_idx), len(lines))
        for line in lines[abstract_idx + 1 : abstract_end]:
            if KEYWORD_RE.match(line) and not keywords:
                keywords = (
                    KEYWORD_RE.sub("", line)
                    .strip()
                    .strip("*")
                    .strip()
                    .lstrip("：:")
                    .strip()
                )
            else:
                abstract_lines.append(line)

    refs: list[tuple[int, str]] = []
    refs_end: int | None = None
    if refs_idx is not None:
        refs_end = next((i for i, _ in headings if i > refs_idx), len(lines))
        current: tuple[int, str] | None = None
        for line in lines[refs_idx + 1 : refs_end]:
            match = REF_DEF_RE.match(line)
            if match:
                if current:
                    refs.append(current)
                current = (int(match.group(1)), match.group(2).strip())
            elif current and line.strip():
                current = (current[0], current[1] + " " + line.strip())
        if current:
            refs.append(current)

    skip: set[int] = set()
    if title_idx is not None:
        skip.add(title_idx)
    if abstract_idx is not None and abstract_end is not None:
        skip.update(range(abstract_idx, abstract_end))

    body_lines: list[str] = []
    for idx, line in enumerate(lines):
        if idx in skip:
            continue
        if refs_idx is not None and refs_end is not None:
            if idx == refs_idx:
                body_lines.append(REFS_PLACEHOLDER)
                continue
            if refs_idx < idx < refs_end:
                continue
        body_lines.append(line)

    # 兜底收集散落在正文里的脚注定义，避免被 pandoc 转成页脚脚注。
    kept_lines: list[str] = []
    for line in body_lines:
        match = REF_DEF_RE.match(line)
        if match:
            number = int(match.group(1))
            if all(number != existing for existing, _ in refs):
                refs.append((number, match.group(2).strip()))
        else:
            kept_lines.append(line)
    refs = sorted({number: content for number, content in refs}.items())

    normalized: list[str] = []
    for line in kept_lines:
        match = MD_HEADING_LINE_RE.match(line)
        if match:
            line = (
                _normalize_heading_level(match.group(1), match.group(2))
                + " "
                + match.group(2).strip()
            )
        normalized.append(line)

    return (
        title,
        "\n".join(abstract_lines).strip(),
        keywords,
        refs,
        "\n".join(normalized),
    )


def _normalize_unicode_greek(markdown: str) -> str:
    """Use explicit math delimiters; dollar closers before digits are ambiguous."""
    parts = MATH_SPAN_RE.split(markdown)
    for index in range(0, len(parts), 2):
        parts[index] = GREEK_CHAR_RE.sub(
            lambda match: "`\\(" + GREEK_UNICODE_MAP[match.group(0)] + "\\)`{=latex}",
            parts[index],
        )
        # Enumeration symbols are often absent from the Latin body font.
        parts[index] = re.sub(
            "[①-⑳]", lambda m: f"({ord(m[0]) - ord('①') + 1})", parts[index]
        )
    return "".join(parts)


def _markdown_fragment_to_latex(markdown: str, resource_path: Path) -> str:
    """pandoc 转正文片段；[^n] 引用替换为上标，行尾统一为 LF。

    pandoc 在 Windows 返回 CRLF，文本模式写文件会二次转换成 \\r\\r\\n，
    多余空行会让 longtable 列声明解析失败。
    """
    markdown = _normalize_unicode_greek(markdown)
    markdown = FOOTNOTE_REF_RE.sub(
        lambda match: f"`\\textsuperscript{{[{match.group(1)}]}}`{{=latex}}", markdown
    )
    try:
        fragment = convert_text(
            markdown,
            to="latex",
            format="markdown-auto_identifiers+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_html",
            extra_args=[
                f"--resource-path={resource_path}",
                "--wrap=none",
                # Supported by both older system Pandoc and the bundled version.
                "--no-highlight",
            ],
        )
    except WorkCancelled:
        raise
    except Exception as exc:
        raise PaperRenderError(f"Pandoc 生成 LaTeX 失败: {exc}") from exc
    fragment = _normalize_vector_math(
        fragment.replace("\r\n", "\n").replace("\r", "\n")
    )
    return _normalize_inline_literals(fragment)


def _normalize_inline_literals(fragment: str) -> str:
    """Wrap inline identifiers at separators; keep full numeric values intact."""

    def render(match: re.Match[str]) -> str:
        if match[1] is None:
            return match[0]  # Never rewrite a verbatim source listing.
        value = match[1]
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", value):
            # Numeric quantities use math rather than the wider code font.
            return (
                r"\("
                + re.sub(r"[eE]", lambda e: r"\mathrm{" + e[0] + "}", value)
                + r"\)"
            )
        if re.fullmatch(r"[0-9a-fA-F]{12,}", value) and re.search(r"[a-fA-F]", value):
            # Hash prefixes have no separators, but must remain copyable in full.
            return (
                r"\texttt{"
                + r"\allowbreak{}".join(
                    value[start : start + 4] for start in range(0, len(value), 4)
                )
                + "}"
            )

        def break_after(token: re.Match[str]) -> str:
            # Consume TeX escapes as whole tokens: punctuation in command names
            # or empty arguments must never become a break inside that command.
            if token[0] in {
                r"\_",
                r"\{",
                r"\}",
                ".",
                "/",
                "=",
                ",",
                ";",
                ":",
                "(",
                ")",
                "[",
                "]",
            } and not value.startswith(r"\allowbreak{}", token.end()):
                return token[0] + r"\allowbreak{}"
            return token[0]

        return (
            r"\texttt{"
            + re.sub(
                r"\\[A-Za-z]+(?:\{\})?|\\.(?:\{\})?|[./=,;:()\[\]]",
                break_after,
                value,
            )
            + "}"
        )

    fragment = re.sub(
        # Pandoc emits quotes as \textquotesingle{} and dictionary braces as
        # \{ / \}. Match these atoms too, while leaving arbitrary nested TeX alone.
        r"\\begin\{verbatim\}.*?\\end\{verbatim\}|"
        r"\\texttt\{((?:\\(?:[A-Za-z]+\{\}|[^A-Za-z](?:\{\})?)|[^\\{}])*)\}",
        render,
        fragment,
        flags=re.S,
    )
    # Permit a break between a math label (e.g. beta=) and its numeric value.
    return re.sub(
        r"\\begin\{verbatim\}.*?\\end\{verbatim\}|\\\)\\\(",
        lambda m: m[0] if m[0].startswith(r"\begin") else r"\)\allowbreak{}\(",
        fragment,
        flags=re.S,
    )


def _normalize_vector_math(latex: str) -> str:
    """Use bold italic math for vector glyphs, without changing operators or text."""
    return re.sub(
        r"\\\(.*?\\\)|\\\[.*?\\\]",
        lambda span: re.sub(
            r"\\(?:mathbf|boldsymbol)\s*\{", lambda _: r"\bm{", span[0]
        ),
        latex,
        flags=re.S,
    )


def _number_display_equations(latex: str) -> str:
    """单行展示公式转为 equation 环境编号（右对齐），多行裸公式保持原样。"""

    def repl(match: re.Match[str]) -> str:
        content = match.group(1)
        if "\\\\" in content and "\\begin{" not in content:
            return match.group(0)
        # Split long lists of independent equations at explicit top-level
        # spacing, never inside fractions, text, cases or existing alignments.
        if len(content) > 180 and "\\begin{" not in content:
            cuts = []
            depth = 0
            for token in re.finditer(r"\\[{}]|[{}]|[,;]\s*\\qquad\b", content):
                value = token.group()
                if value == "{":
                    depth += 1
                elif value == "}":
                    depth -= 1
                elif depth == 0 and value[0] in ",;":
                    cuts.append((token.start() + 1, token.end()))
            if cuts:
                rows = []
                start = 0
                for end, following in cuts:
                    rows.append(content[start:end].strip())
                    start = following
                rows.append(content[start:].strip())
                content = (
                    "\\begin{gathered}\n" + " \\\\\n".join(rows) + "\n\\end{gathered}"
                )
        return "\\begin{equation}" + content.strip() + "\\end{equation}"

    return DISPLAY_MATH_RE.sub(repl, latex)


def build_china_paper_preamble() -> str:
    """国赛风格版式：宋体正文、单倍行距、黑体标题、2.5cm 页边距、页脚居中页码。

    标题字号对齐优秀论文惯例：一级标题四号黑体居中，二/三级小四黑体左对齐；
    作者手写编号（一、/1.1/5.1.1）直接作为标题文字，故关闭自动编号。
    """
    return r"""% !TEX program = xelatex
% Remit-LaTeX-Assembler: china-v2
\documentclass[UTF8,a4paper,zihao=-4,linespread=1.08]{ctexart}
\usepackage[top=2.5cm,bottom=2.5cm,left=2.5cm,right=2.5cm]{geometry}
\usepackage{amsmath,amssymb,bm}
\usepackage{graphicx}
% 允许行内公式在逗号后断行，避免长数值序列顶出版心
\makeatletter
\mathchardef\remit@origcomma=\mathcode`,
\mathcode`,="8000
\begingroup
\catcode`,=\active
\gdef,{\remit@origcomma\penalty0\relax}
\endgroup
\makeatother
\usepackage{longtable,booktabs,array}
\usepackage{caption}
\captionsetup{labelsep=quad}
\captionsetup[table]{skip=6pt}
\captionsetup[figure]{skip=6pt}
\newcounter{none} % for unnumbered tables
\usepackage{calc}
\usepackage{etoolbox}
\makeatletter
\patchcmd\longtable{\par}{\if@noskipsec\mbox{}\fi\par}{}{}
% longtable continuation boxes can be split by the current LaTeX output routine.
% Stretch-only glue preserves bottom alignment without infinite shrink errors.
\patchcmd\LT@output{\vss}{\vfil}{}{}
\patchcmd\LT@output{\vss}{\vfil}{}{}
\makeatother
\IfFileExists{footnotehyper.sty}{\usepackage{footnotehyper}}{\usepackage{footnote}}
\makesavenoteenv{longtable}
\usepackage{float}
\usepackage{indentfirst}
\setlength{\parindent}{2em}
\setlength{\emergencystretch}{3em}
\raggedbottom
% Let following text fill space when a figure cannot fit at its callout.
\renewcommand{\topfraction}{0.68}
\renewcommand{\bottomfraction}{0.5}
\renewcommand{\textfraction}{0.2}
\renewcommand{\floatpagefraction}{0.9}
\setcounter{topnumber}{1}
\setcounter{bottomnumber}{1}
\setcounter{totalnumber}{2}
\setlength{\textfloatsep}{12pt plus 2pt minus 2pt}
\setlength{\intextsep}{10pt plus 2pt minus 2pt}
\makeatletter
\newsavebox\pandoc@box
% 常规插图同宽，只有极高的纵向图才触发安全高度限制。
\newcommand*\pandocbounded[1]{%
  \sbox\pandoc@box{#1}%
  \Gscale@div\@tempa{0.68\linewidth}{\wd\pandoc@box}%
  \Gscale@div\@tempb{0.42\textheight}{\dimexpr\ht\pandoc@box+\dp\pandoc@box\relax}%
  \ifdim\@tempb\p@<\@tempa\p@\let\@tempa\@tempb\fi%
  \scalebox{\@tempa}{\usebox\pandoc@box}%
}
\def\fps@figure{htb}
\makeatother
\providecommand{\tightlist}{\setlength{\itemsep}{0pt}\setlength{\parskip}{0pt}}
\usepackage{fvextra}
\RecustomVerbatimEnvironment{verbatim}{Verbatim}{fontsize=\small,breaklines=true,breakanywhere=true}
\ctexset{
  section = {numbering=false, format=\centering\heiti\zihao{4}},
  subsection = {numbering=false, format=\heiti\zihao{-4}},
  subsubsection = {numbering=false, format=\heiti\zihao{-4}},
  paragraph = {numbering=false, format=\heiti\zihao{-4}},
}
\usepackage{fancyhdr}
\pagestyle{fancy}
\fancyhf{}
\fancyfoot[C]{\small\thepage}
\renewcommand{\headrulewidth}{0pt}
\fancypagestyle{plain}{\fancyhf{}\fancyfoot[C]{\small\thepage}\renewcommand{\headrulewidth}{0pt}}
\usepackage[hidelinks]{hyperref}
"""


def _build_references_latex(refs: list[tuple[int, str]]) -> str:
    """编号著录的参考文献节；条目为模型给出的纯文本，逐字符转义。"""
    if not refs:
        return ""
    items = "\n".join(
        f"\\noindent [{number}] {_escape_latex_text(content)}\\par"
        for number, content in refs
    )
    return (
        "\\section{参考文献}\n\\begingroup\n\\zihao{5}\n"
        "\\setlength{\\parindent}{0pt}\n\\setlength{\\parskip}{3pt}\n"
        f"{items}\n\\endgroup\n"
    )


SYMBOL_HEADING_RE = re.compile(r"(?m)^\\(?:sub)?section\{[^}]*符号说明[^}]*\}\s*$")
SYMBOL_REGION_END_RE = re.compile(r"(?m)^\\(?:sub)?section\{")
LONGTABLE_START_RE = re.compile(
    r"\{\\def\\LTcaptype\{none\}[^\n]*\n"
    r"\\begin\{longtable\}(?:\[[^\]]*\])?\{"
)
LONGTABLE_END_RE = re.compile(r"\\end\{longtable\}\n\}\n?")
MINIPAGE_CELL_RE = re.compile(
    r"\\begin\{minipage\}\[b\]\{\\linewidth\}\\raggedright\s*(.*?)\s*\\end\{minipage\}",
    re.S,
)
SUBSUBSECTION_LINE_RE = re.compile(r"(?m)^\\subsubsection\{[^}]*\}\s*\n?")
SYMBOL_FULL_WIDTH_SPECS = {
    2: r"@{}>{\centering\arraybackslash}p{0.22\linewidth}>{\centering\arraybackslash}p{0.68\linewidth}@{}",
    3: (
        r"@{}>{\centering\arraybackslash}p{0.25\linewidth}"
        r">{\centering\arraybackslash}p{0.54\linewidth}"
        r">{\centering\arraybackslash}p{0.12\linewidth}@{}"
    ),
}


def _match_braces(text: str, open_index: int) -> int:
    depth = 0
    for index in range(open_index, len(text)):
        char = text[index]
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return index
    return -1


def _find_symbol_tables(region: str) -> list[tuple[int, int, str, str]]:
    """找出区域内每个 longtable 块的 (起点, 终点, 列声明, 表体)。

    列声明可能跨行并含嵌套花括号（pandoc 通栏表的 \\real{} 写法），
    用花括号配对而不是单纯正则来截取。
    """
    blocks = []
    for match in LONGTABLE_START_RE.finditer(region):
        open_brace = match.end() - 1
        close_brace = _match_braces(region, open_brace)
        if close_brace < 0:
            continue
        end = LONGTABLE_END_RE.search(region, close_brace + 1)
        if not end:
            continue
        blocks.append(
            (
                match.start(),
                end.end(),
                region[open_brace + 1 : close_brace],
                region[close_brace + 1 : end.start()],
            )
        )
    return blocks


def _table_column_count(spec: str) -> int:
    if "\\real{" in spec:
        return spec.count("\\real{")
    cleaned = re.sub(r"@\{[^{}]*\}", "", spec)
    cleaned = re.sub(r">\{[^{}]*\}", "", cleaned)
    return len(re.findall(r"[lcr]", cleaned))


def _merge_symbol_section_tables(body_tex: str) -> str:
    """把"符号说明"小节里被拆分的多张子表合并成一张通栏三线表。

    写作手容易按类别把符号表拆成 4.1.1/4.1.2 等子表，而比赛惯例是
    一整张通栏三线表。这里在 LaTeX 层面兜底：丢弃子节标题、按出现
    顺序拼接各表数据行，并把列声明统一为通栏 p 列；列数不一致时
    保持原样，避免产出无法编译的表格。
    """
    headings = list(SYMBOL_HEADING_RE.finditer(body_tex))
    if not headings:
        return body_tex
    subsections = [m for m in headings if m.group(0).startswith("\\subsection")]
    heading = subsections[0] if subsections else headings[0]
    boundary = SYMBOL_REGION_END_RE.search(body_tex, heading.end())
    region_end = boundary.start() if boundary else len(body_tex)
    region = body_tex[heading.start() : region_end]
    if not _find_symbol_tables(region):
        return body_tex
    region = SUBSUBSECTION_LINE_RE.sub("", region)
    tables = _find_symbol_tables(region)
    if not tables:
        return body_tex

    first_count = _table_column_count(tables[0][2])
    if any(_table_column_count(table[2]) != first_count for table in tables[1:]):
        return body_tex

    row_chunks = [tables[0][3]]
    for _, _, _, chunk in tables[1:]:
        if "\\endlastfoot" in chunk:
            chunk = chunk.split("\\endlastfoot", 1)[1]
        elif "\\endhead" in chunk:
            chunk = chunk.split("\\endhead", 1)[1]
        row_chunks.append(chunk.strip("\n") + "\n")
    merged_body = MINIPAGE_CELL_RE.sub(lambda m: m.group(1), "".join(row_chunks))
    if "\\endfirsthead" not in merged_body and "\\endhead" in merged_body:
        header, rows = merged_body.split("\\endhead", 1)
        # A caption belongs to the first head only, not every continuation page.
        merged_body = header + "\\endfirsthead\n" + header + "\\endhead\n" + rows
    spec = SYMBOL_FULL_WIDTH_SPECS.get(first_count, tables[0][2])
    # 符号表是论文里少数必须带题注的通栏表：表题置于表上方，整表居中，
    # 不再借用 pandoc 的 LTcaptype=none 包装（那会吞掉表号）。
    merged = (
        f"\\begin{{longtable}}[c]{{{spec}}}\n"
        "\\caption{主要符号说明}\\\\\n"
        + merged_body.strip("\n")
        + "\n\\end{longtable}\n"
    )
    if "\\endlastfoot" in merged_body:
        data_rows = merged_body.split("\\endlastfoot", 1)[1].strip()
        row_count = len(re.findall(r"\\\\\s*(?:\n|$)", data_rows))
        if 0 < row_count <= 18 and len(data_rows) < 5000:
            # The bounded core-symbol table fits on one page. Keep its rows
            # together instead of leaving a few at the bottom of the prior page.
            header = re.split(r"\\endfirsthead|\\endhead", merged_body, maxsplit=1)[0]
            merged = (
                "\\begin{table}[!htbp]\n\\centering\n\\caption{主要符号说明}\n"
                + f"\\begin{{tabular}}{{{spec}}}\n"
                + header.strip()
                + "\n"
                + data_rows
                + "\n\\bottomrule\n\\end{tabular}\n\\end{table}\n"
            )

    span_start, span_end = tables[0][0], tables[-1][1]
    between = region[span_start:span_end]
    leftover_parts = []
    cursor = 0
    for start, end, _, _ in tables:
        leftover_parts.append(between[cursor : start - span_start])
        cursor = end - span_start
    leftover_parts.append(between[cursor:])
    leftover = "".join(leftover_parts).strip()
    replacement = (leftover + "\n\n" if leftover else "") + merged
    region = region[:span_start] + replacement + region[span_end:]
    return body_tex[: heading.start()] + region + body_tex[region_end:]


INLINE_IMAGE_SPLIT_RE = re.compile(r"(!\[[^\]]*\]\([^)]+\))")
BRACE_FOOTNOTE_DEF_RE = re.compile(r"\{\^?\[(\d+)\]:\s*([^{}\n]+?)\}")
FIGURE_ENV_RE = re.compile(r"\\begin\{figure\}.*?\\end\{figure\}", re.S)
FILENAME_CAPTION_RE = re.compile(
    r"\\caption\{[^{}]*\.(?:png|jpg|jpeg|pdf|svg)[^{}]*\}\s*\n?", re.I
)
FIGURE_OR_HEADING_RE = re.compile(
    r"\\(?:sub)?section\{[^{}]*\}|\\begin\{figure\}.*?\\end\{figure\}", re.S
)
FIGURE_GRAPHICS_RE = re.compile(r"\\includegraphics(?:\[[^\]]*\])?\{([^{}]*)\}")
FIGURE_CAPTION_RE = re.compile(r"\\caption\{(?P<text>.*?)\}\s*", re.S)
MANUAL_FIGNO_RE = re.compile(
    r"^\s*(?:图|Fig(?:ure)?\.?)\s*(\d+(?:[-–.]\d+)?(?:[（(][a-zA-Z][）)])?)\s*[:：．.、]?\s*"
)
SUBFIG_TAG_RE = re.compile(r"^\s*[（(][a-zA-Z][）)]\s*")
_CN_SECTION_NUMBERS = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
    "十一": 11,
    "十二": 12,
}


def _convert_brace_footnotes(markdown: str) -> str:
    """把写作手误写成 {^[n]: 著录} 或 {[^n]: 著录} 的内嵌文献还原成标准脚注。

    原地留下 [^n] 引用，著录文本移入文末参考文献区，交给既有编号著录
    流程；同编号已有定义时丢弃重复著录。
    """
    collected: list[tuple[str, str]] = []

    def repl(match: re.Match[str]) -> str:
        number, text = match.group(1), match.group(2).strip()
        collected.append((number, text))
        return f"[^{number}]"

    converted = BRACE_FOOTNOTE_DEF_RE.sub(repl, markdown)
    additions = [
        f"[^{number}]: {text}"
        for number, text in collected
        if f"[^{number}]:" not in converted
    ]
    if not additions:
        return converted
    lines = converted.splitlines()
    refs_idx = next(
        (idx for idx, line in enumerate(lines) if REFS_HEADING_RE.match(line)),
        None,
    )
    insert_at = refs_idx + 1 if refs_idx is not None else len(lines)
    lines[insert_at:insert_at] = [""] + additions + [""]
    return "\n".join(lines)


def _extract_inline_images(markdown: str) -> str:
    """把混在正文行内的图片提取为独立图片段落。

    图片与文字同段时 pandoc 会把图片当作行内元素，导致排版错乱且没有
    题注；独立成段后由 implicit_figures 生成居中 figure 与题注。
    """
    out: list[str] = []
    for line in markdown.splitlines():
        stripped = line.lstrip()
        if "![" not in line or stripped.startswith("|") or stripped.startswith("#"):
            out.append(line)
            continue
        pending: list[str] = []
        for part in INLINE_IMAGE_SPLIT_RE.split(line):
            part = part.strip()
            if not part:
                continue
            if IMAGE_BLOCK_RE.match(part):
                if pending:
                    out.append("".join(pending))
                    pending = []
                out.extend(["", part, ""])
            else:
                pending.append(part)
        if pending:
            out.append("".join(pending))
    return "\n".join(out)


def _center_longtables(body_tex: str) -> str:
    """pandoc 生成的 longtable 默认贴左，比赛惯例是三线表整体居中。"""
    body_tex = body_tex.replace("\\begin{longtable}[]", "\\begin{longtable}[c]")
    body_tex = body_tex.replace(
        r"\raggedright\arraybackslash", r"\centering\arraybackslash"
    )
    body_tex = body_tex.replace(r"\raggedright", r"\centering")
    # Compact Pandoc column specs (l/r/c) need centering as well.
    return re.sub(
        r"(\\begin\{longtable\}\[c\]\{)([@{}lrc ]+)(\})",
        lambda m: m[1] + re.sub("[lr]", "c", m[2]) + m[3],
        body_tex,
    )


def _normalize_numeric_tables(body_tex: str) -> str:
    """Give short numeric tables natural column widths and explicit unit rows.

    Long/text-heavy tables retain their multipage layout. Cell values are kept
    byte-for-byte; only header layout and the surrounding environment change.
    """
    for start, end, spec, body in reversed(_find_symbol_tables(body_tex)):
        if "\\endlastfoot" not in body:
            continue
        data = body.split("\\endlastfoot", 1)[1].strip()
        rows = re.split(r"\\\\\s*(?:\n|$)", data)
        rows = [row.strip() for row in rows if row.strip()]
        count = _table_column_count(spec)
        if not 2 <= count <= 9 or not 1 <= len(rows) <= 24:
            continue
        cells = [cell.strip() for row in rows for cell in row.split("&")]
        numeric = sum(
            bool(re.fullmatch(r"[-+\d.,eE%\\(){} ]+|是|否|—|-", cell)) for cell in cells
        )
        compact_text_table = (
            count <= 5
            and "\\real{" not in spec
            and all(len(cell) <= 60 for cell in cells)
        )
        if numeric < len(cells) * 0.7 and not compact_text_table:
            continue
        header_match = re.search(
            r"\\toprule(?:\\noalign\{\})?\s*(.*?)\\midrule", body, re.S
        )
        if not header_match:
            continue
        header = MINIPAGE_CELL_RE.sub(
            lambda m: m.group(1).strip(), header_match.group(1)
        ).strip()
        header = re.sub(r"\\\\\s*$", "", header)
        headers = header.split("&")
        if len(headers) != count:
            continue
        formatted = [" ".join(cell.split()) for cell in headers]
        title = re.search(
            r"(?:^|\n\n)(表\s*\d+(?:[-–.]\d+)?[　 \t]+[^\n]{1,120})\s*\n\n$",
            body_tex[:start],
        )
        caption = ""
        if title and "。" not in title[1]:
            caption = "\\caption*{" + title[1].strip() + "}\n"
            start = title.start()
        parsed_rows = [re.split(r"(?<!\\)&", row) for row in rows]
        if any(len(row) != count for row in parsed_rows):
            continue
        # Natural-width columns do not shrink to tabular*'s requested width.
        # Split wide numeric tables, repeating their row key, instead of scaling
        # the font or rounding any scientific value. The PDF gate still checks
        # actual geometry: this conservative estimate is not an acceptance test.
        widths = []
        for col in range(count):
            values = [formatted[col], *(row[col].strip() for row in parsed_rows)]
            widths.append(max(_table_text_width(value) for value in values) + 2)
        groups = [list(range(count))]
        if sum(widths) > 72:
            groups = []
            group = [0]
            for col in range(1, count):
                if len(group) > 1 and sum(widths[i] for i in group) + widths[col] > 72:
                    groups.append(group)
                    group = [0]
                group.append(col)
            groups.append(group)
        tables = []
        for group_index, group in enumerate(groups):
            group_caption = caption
            if len(groups) > 1:
                group_caption += f"\\caption*{{列分组 {group_index + 1}/{len(groups)}（首列对应同一行）}}\n"
            group_data = (
                " \\\\\n".join(
                    " & ".join(row[col].strip() for col in group) for row in parsed_rows
                )
                + r" \\"
            )
            tables.append(
                "\\begin{table}[!htbp]\n\\centering\n\\small\n"
                + group_caption
                + "\\begin{tabular*}{\\linewidth}{@{\\extracolsep{\\fill}}"
                + "c" * len(group)
                + "@{}}\n\\toprule\n"
                + " & ".join(formatted[col] for col in group)
                + " \\\\\n\\midrule\n"
                + group_data
                + "\n\\bottomrule\n\\end{tabular*}\n\\end{table}\n"
            )
        replacement = "\n".join(tables)
        body_tex = body_tex[:start] + replacement + body_tex[end:]
    return body_tex


def _table_text_width(value: str) -> int:
    """Estimate a cell width in Latin characters, allowing two for CJK."""
    visible = re.sub(r"\\[a-zA-Z]+\s*|[{}]", "", value)
    visible = visible.replace(r"\_", "_")
    return sum(2 if "\u4e00" <= char <= "\u9fff" else 1 for char in visible)


def _china_section_number(title: str) -> int | None:
    """从手写标题推断一级节号："五、…" 或 "5.1 …" 两种写法。"""
    match = re.match(rf"^\s*([{CN_NUM_CHARS}]{{1,3}})\s*[、.．]", title)
    if match:
        return _CN_SECTION_NUMBERS.get(match.group(1))
    match = re.match(r"^\s*(\d+)(?:\.\d+)*(?=\D|$)", title)
    if match:
        return int(match.group(1))
    return None


def _normalize_figures(body_tex: str) -> str:
    """统一 figure 排版：固定图幅、独立成图、题注规范为"图 N 描述"。

    写作手给的 alt 可能带手写编号（"图 6-3 …"）、文件名或 (a) 子图
    标记；题注统一用不自动编号的 \\caption*，编号沿用作者手写值，
    缺省时按当前一级节自动编排（图 5-1、图 6-2…），与正文引用一致。
    """
    tokens = list(FIGURE_OR_HEADING_RE.finditer(body_tex))
    if not any(token.group(0).startswith("\\begin{figure}") for token in tokens):
        return body_tex

    section = 0
    counter = 0
    pieces: list[str] = []
    cursor = 0
    for token in tokens:
        text = token.group(0)
        if not text.startswith("\\begin{figure}"):
            title = text[text.index("{") + 1 : -1]
            number = _china_section_number(title)
            if number is not None and (text.startswith("\\section{") or section == 0):
                section = number
                counter = 0
            continue

        pieces.append(body_tex[cursor : token.start()])
        cursor = token.end()
        block = FILENAME_CAPTION_RE.sub("", text)
        graphics_match = FIGURE_GRAPHICS_RE.search(block)
        if not graphics_match:
            pieces.append(block)
            continue
        caption_text = ""
        caption_match = re.search(r"\\caption\{", block)
        if caption_match:
            close = _match_braces(block, caption_match.end() - 1)
            if close >= 0:
                caption_text = " ".join(block[caption_match.end() : close].split())
        # Some writers put the real caption in a separate paragraph after an
        # image whose alt text is just its filename. Move that title into the
        # figure so it cannot be duplicated or stranded by float placement.
        following = re.match(
            r"\s*\n(图\s*\d+(?:[-–.]\d+)?(?:[（(][a-zA-Z][）)])?[　 \t]+[^\n]{1,160})\n",
            body_tex[cursor:],
        )
        if (
            not caption_text
            and following
            and "。" not in following[1]
            and not re.match(
                r"图\s*\d+(?:[-–.]\d+)?\s+(?:给出|显示|表明|可见|中)", following[1]
            )
        ):
            caption_text = following[1].strip()
            cursor += following.end()
        manual = MANUAL_FIGNO_RE.match(caption_text)
        if manual:
            label = f"图 {manual.group(1)}"
            caption_text = MANUAL_FIGNO_RE.sub("", caption_text)
            number = re.fullmatch(rf"{section}[-–.](\d+)", manual.group(1))
            if number:
                counter = max(counter, int(number[1]))
        else:
            counter += 1
            label = f"图 {section}-{counter}" if section else f"图 {counter}"
        caption_text = SUBFIG_TAG_RE.sub("", caption_text).strip(" ，,。；;")
        caption = f"{label}　{caption_text}" if caption_text else label
        pieces.append(
            "\\begin{figure}[htb]\n\\centering\n"
            "\\includegraphics[width=0.68\\linewidth,"
            "height=0.42\\textheight,keepaspectratio]"
            f"{{{graphics_match.group(1)}}}\n"
            f"\\caption*{{{caption}}}\n\\end{{figure}}\n"
        )
    pieces.append(body_tex[cursor:])
    return "".join(pieces)


def _resolve_figure_callouts(body_tex: str) -> str:
    """Resolve file-based citations only after actual figure selection/numbering."""
    labels = {}
    graphics = {}
    prose = FIGURE_ENV_RE.sub("", body_tex)

    def reuse_figure(match: re.Match[str]) -> str:
        block = match[0]
        graphic = FIGURE_GRAPHICS_RE.search(block)
        caption = re.search(
            r"\\caption\*?\{(图\s*\d+(?:[-–.]\d+)?(?:[（(][a-zA-Z][）)])?)", block
        )
        if graphic and caption:
            name = Path(graphic[1]).name
            if name in labels:
                if graphics[name] != graphic[0]:
                    raise PaperRenderError(f"图片重复插入，图号无法唯一确定：{name}")
                if caption[1] != labels[name] and re.search(
                    re.escape(caption[1]).replace(r"\ ", r"\s*") + r"(?![\d.−–-])",
                    prose,
                ):
                    raise PaperRenderError(
                        f"重复图片 {name} 还有手写图号引用，请改为 @fig:{name}@ 后重试"
                    )
                # The same asset may support multiple chapters. Render it once
                # and retain the later caption's explanation beside a callout.
                opening = block.index("{", caption.start())
                closing = _match_braces(block, opening)
                if closing < 0:
                    raise PaperRenderError(f"图片题注不完整：{name}")
                explanation = block[caption.end() : closing].strip()
                return labels[name] + "（前文已列图）：" + explanation + "\n"
            labels[name] = caption[1]
            graphics[name] = graphic[0]
        return block

    body_tex = FIGURE_ENV_RE.sub(reuse_figure, body_tex)

    def resolve(match):
        name = Path(match[1].replace(r"\_", "_")).name
        if name not in labels:
            raise PaperRenderError(f"正文引用了未插入的图片：{name}")
        return labels[name]

    return re.sub(r"(?:图\s*)?@fig:([^@\n]+)@", resolve, body_tex)


def _assemble_china_paper(
    markdown: str, output_path: Path, resource_path: Path, *, mode: str = "full_paper"
) -> None:
    """中文赛事论文：摘要专用页 + 正文片段 + 参考文献，组装为完整 LaTeX。

    摘要页按国赛惯例组织：三号黑体居中标题、居中"摘 要"、小四正文、
    黑体"关键词："引导，页码从摘要页开始；参考文献保持题目要求的
    位置（正文之后、附录之前），由占位符回插。
    """
    markdown = _convert_brace_footnotes(markdown)
    title, abstract_md, keywords, refs, body_md = _split_china_paper(markdown)
    if mode == "short_report" and "短报告" not in title:
        title = (title or "计算结果") + "（短报告）"
    body_md = _extract_inline_images(body_md)
    parts = [build_china_paper_preamble(), "\\begin{document}\n"]
    title_tex = _escape_latex_text(title) if title else "数学建模论文"
    if abstract_md:
        parts.append(
            "\\begin{center}\n{\\heiti\\zihao{3} "
            + title_tex
            + "}\n\\par\\vspace{1.5em}\n{\\heiti\\zihao{4} 摘\\hspace{0.5em}要}\n"
            "\\end{center}\n\\vspace{0.8em}\n"
        )
        # A center environment suppresses indentation of the following paragraph.
        # Explicit horizontal space is stable across ctex/platform font choices.
        parts.append(
            "\\noindent\\hspace*{2em}%\n"
            + _markdown_fragment_to_latex(abstract_md, resource_path).lstrip()
        )
        if keywords:
            parts.append(
                "\\par\\vspace{1em}\n\\noindent{\\heiti\\bfseries 关键词："
                + _escape_latex_text(_normalize_keywords(keywords))
                + "}\n"
            )
        if mode != "short_report":
            parts.append("\\newpage\n")
    else:
        parts.append(
            "\\begin{center}{\\heiti\\zihao{3} " + title_tex + "}\\end{center}\n"
        )
    body_tex = _number_display_equations(
        _markdown_fragment_to_latex(body_md, resource_path)
    )
    body_tex = _merge_symbol_section_tables(body_tex)
    body_tex = _normalize_numeric_tables(body_tex)
    body_tex = _resolve_figure_callouts(
        _normalize_figures(_center_longtables(body_tex))
    )
    refs_tex = _build_references_latex(refs)
    if REFS_PLACEHOLDER in body_tex:
        body_tex = body_tex.replace(
            REFS_PLACEHOLDER, "\n" + refs_tex if refs_tex else ""
        )
    elif refs_tex:
        body_tex += "\n" + refs_tex
    parts.append(body_tex)
    parts.append("\n\\end{document}\n")
    output_path.write_text("".join(parts), encoding="utf-8", newline="\n")


def _validate_latex_source(tex_path: Path) -> None:
    if not tex_path.is_file() or tex_path.stat().st_size < 500:
        raise PaperRenderError("生成的 res.tex 为空或内容异常短")
    source = tex_path.read_text(encoding="utf-8")
    required = ("\\documentclass", "\\begin{document}", "\\end{document}")
    missing = [token for token in required if token not in source]
    if missing:
        raise PaperRenderError("res.tex 缺少完整文档结构: " + ", ".join(missing))
    # 附录源码里的注释行（# ...）不是 Markdown 泄漏，先剔除逐字代码块再扫描。
    scan_source = VERBATIM_BLOCK_RE.sub("", source)
    markdown_leaks = re.findall(
        r"(?m)^\s*(?:#{1,6}\s+|!\[[^\]]*\]\([^)]+\))", scan_source
    )
    if markdown_leaks:
        raise PaperRenderError("res.tex 仍包含未转换的 Markdown 块标记")


def _compile_latex(
    tex_path: Path, build_dir: Path, *, resource_path: Path | None = None
) -> str:
    configured = settings.LATEX_ENGINE.strip() or "xelatex"
    engine = shutil.which(configured)
    if not engine:
        raise PaperRenderError(
            f"找不到 LaTeX 编译器 {configured}；请安装 MiKTeX 或 TeX Live 并加入 PATH"
        )

    command = [
        engine,
        "-no-shell-escape",
        "-interaction=nonstopmode",
        "-halt-on-error",
        "-file-line-error",
        f"-output-directory={build_dir}",
        str(tex_path.resolve()) if resource_path is not None else tex_path.name,
    ]
    logs: list[str] = []
    for compile_pass in range(1, 3):
        try:
            completed = run_xelatex(
                command,
                pdf_path=build_dir / (tex_path.stem + ".pdf"),
                cwd=resource_path or tex_path.parent,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=settings.LATEX_COMPILE_TIMEOUT_SECONDS,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise PaperRenderError(
                f"LaTeX 第 {compile_pass} 次编译超过 "
                f"{settings.LATEX_COMPILE_TIMEOUT_SECONDS:g} 秒"
            ) from exc
        output = (completed.stdout or "") + "\n" + (completed.stderr or "")
        logs.append(f"===== pass {compile_pass} =====\n{output}")
        if completed.returncode != 0:
            (build_dir / "compile-output.txt").write_text(
                "\n".join(logs), encoding="utf-8"
            )
            tail = output[-3000:].strip()
            raise PaperRenderError(
                f"res.tex 第 {compile_pass} 次编译失败（退出码 "
                f"{completed.returncode}）：{tail}"
            )
    (build_dir / "compile-output.txt").write_text("\n".join(logs), encoding="utf-8")
    return Path(engine).name


def inspect_pdf_artifact(
    pdf_path: str | Path,
    *,
    comp_template: CompTemplate,
    minimum_pages: int,
) -> dict[str, object]:
    """重开并渲染抽样页面，拒绝空白、损坏或纸型错误的 PDF。"""
    import pymupdf

    path = Path(pdf_path)
    if not path.is_file() or path.stat().st_size < 1_000:
        raise PaperRenderError("res.pdf 缺失或内容异常短")
    try:
        document = pymupdf.open(path)
    except Exception as exc:
        raise PaperRenderError(f"res.pdf 无法重新打开: {exc}") from exc

    try:
        if document.needs_pass:
            raise PaperRenderError("res.pdf 被意外加密，无法验收")
        page_count = document.page_count
        if page_count < minimum_pages:
            raise PaperRenderError(
                f"res.pdf 仅 {page_count} 页，少于终稿下限 {minimum_pages} 页"
            )
        first_rect = document[0].rect
        expected = (
            (595.3, 841.9) if comp_template == CompTemplate.CHINA else (612.0, 792.0)
        )
        if (
            abs(first_rect.width - expected[0]) > 8
            or abs(first_rect.height - expected[1]) > 8
        ):
            label = "A4" if comp_template == CompTemplate.CHINA else "US Letter"
            raise PaperRenderError(f"res.pdf 首页纸型不是要求的 {label}")

        blank_pages: list[int] = []
        total_text_chars = 0
        for index, page in enumerate(document):
            page_text = page.get_text().strip()
            total_text_chars += len(page_text)
            if not page_text and not page.get_images() and not page.get_drawings():
                blank_pages.append(index + 1)
            for block in page.get_text("blocks", clip=pymupdf.INFINITE_RECT()):
                x0, y0, x1, y1 = block[:4]
                if (
                    x0 < -2
                    or y0 < -2
                    or x1 > page.rect.width + 2
                    or y1 > page.rect.height + 2
                ):
                    raise PaperRenderError(f"res.pdf 第 {index + 1} 页存在越界文本")
        if blank_pages:
            raise PaperRenderError(
                "res.pdf 含完全空白页: " + ", ".join(map(str, blank_pages))
            )
        if total_text_chars < 1_000:
            raise PaperRenderError(
                "res.pdf 可提取正文不足 1000 字符，疑似字体或渲染损坏"
            )

        sampled_pages = sorted({0, page_count // 2, page_count - 1})
        for index in sampled_pages:
            pixmap = document[index].get_pixmap(
                matrix=pymupdf.Matrix(1, 1), alpha=False
            )
            sample = pixmap.samples[:: max(1, pixmap.n * 20)]
            if not sample or all(value > 248 for value in sample):
                raise PaperRenderError(f"res.pdf 第 {index + 1} 页渲染结果近似全白")
        return {
            "page_count": page_count,
            "paper_size": "A4" if comp_template == CompTemplate.CHINA else "Letter",
            "blank_pages": blank_pages,
            "extracted_text_chars": total_text_chars,
            "rendered_pages_checked": [index + 1 for index in sampled_pages],
        }
    finally:
        document.close()


def _remove_legacy_paper_outputs(root: Path) -> None:
    for filename in _LEGACY_PAPER_OUTPUTS:
        (root / filename).unlink(missing_ok=True)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def render_paper_docx(task_id: str) -> Path:
    """Render polished DOCX for a completed task."""
    work_dir = Path("project") / "work_dir" / task_id
    work_dir.mkdir(parents=True, exist_ok=True)

    source_md = work_dir / "res.md"
    if not source_md.exists():
        raise FileNotFoundError(f"missing paper markdown: {source_md}")

    polished_md = work_dir / "res_polished.md"
    raw_markdown = source_md.read_text(encoding="utf-8")
    polished_markdown = polish_markdown(raw_markdown, work_dir)
    polished_md.write_text(polished_markdown, encoding="utf-8")

    reference_docx = build_reference_docx(work_dir)
    docx_path = work_dir / "res.docx"
    polished_docx_path = work_dir / "res_polished.docx"

    convert_markdown_to_docx(
        markdown=polished_markdown,
        output_path=docx_path,
        reference_docx=reference_docx,
        resource_path=work_dir,
    )
    shutil.copy2(docx_path, polished_docx_path)
    try:
        convert_markdown_to_pdf(
            markdown=polished_markdown,
            output_path=work_dir / "res_polished.pdf",
            resource_path=work_dir,
        )
    except WorkCancelled:
        raise
    except Exception as exc:  # pragma: no cover - best effort export
        logger.warning("pdf export skipped: %s", exc)
    logger.info("paper docx generated: %s", docx_path)
    logger.info("paper polished docx copied: %s", polished_docx_path)
    return docx_path


def polish_markdown(markdown: str, work_dir: Path, *, mode: str = "full_paper") -> str:
    """Apply paper-level Markdown cleanup."""
    markdown = markdown.replace("\r\n", "\n")
    markdown = normalize_common_math(markdown)
    markdown = compact_abstract(markdown)
    markdown = merge_image_blocks(markdown, work_dir)
    if mode != "short_report":
        markdown = append_source_code_appendix(markdown, work_dir)
    markdown = ensure_blank_lines_around_headings(markdown)
    markdown = enforce_em_dash_budget(markdown)
    return markdown.strip() + "\n"


def enforce_em_dash_budget(markdown: str, budget: int = 2) -> str:
    """Keep the final prose within the style auditor's em-dash budget.

    Writer output can contain a few rhetorical em dashes even when prompted not
    to.  Preserve the first occurrences for intentional emphasis and normalize
    any excess to a Chinese comma so the final deterministic audit does not
    repeatedly suspend an otherwise complete paper for the same style issue.
    """
    if budget < 0:
        raise ValueError("破折号预算不能为负数")
    seen = 0
    chars: list[str] = []
    for char in markdown:
        if char == "—":
            seen += 1
            chars.append(char if seen <= budget else "，")
        else:
            chars.append(char)
    return "".join(chars)


def normalize_common_math(markdown: str) -> str:
    """Normalize a few math-like tokens so pandoc renders them more reliably."""
    replacements = {
        r"(?<![\\$])\bR\^2\b(?![\\$])": r"$R^2$",
        r"(?<![\\$])\bR\^3\b(?![\\$])": r"$R^3$",
        r"(?<![\\$])\bOOF\s+R\^2\b(?![\\$])": r"OOF $R^2$",
    }
    lines = markdown.splitlines()
    out: list[str] = []
    in_display_math = False

    idx = 0
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        if stripped == "$$":
            out.append(line)
            in_display_math = not in_display_math
            idx += 1
            continue
        if in_display_math:
            if stripped == "R^2" and idx + 1 < len(lines):
                next_stripped = lines[idx + 1].strip()
                if next_stripped.startswith("="):
                    out.append("R^2 " + next_stripped)
                    idx += 2
                    continue
            out.append(line)
            idx += 1
            continue
        for pattern, repl in replacements.items():
            line = re.sub(pattern, repl, line)
        out.append(line)
        idx += 1

    return "\n".join(out)


def compact_abstract(markdown: str) -> str:
    """Normalize keywords without deleting the writer's evidence or emphasis.

    Length corrections belong to the writer's bounded revision loop. String
    truncation here used to remove results and silently undo bold formatting.
    """
    lines = markdown.splitlines()
    start = None
    start_level = 0

    for idx, line in enumerate(lines):
        if ABSTRACT_RE.match(line):
            start = idx
            start_level = len(line) - len(line.lstrip("#"))
            break
    if start is None:
        return markdown

    end = len(lines)
    for idx in range(start + 1, len(lines)):
        line = lines[idx]
        if HEADING_RE.match(line):
            level = len(line) - len(line.lstrip("#"))
            if level <= start_level:
                end = idx
                break

    prefix = lines[: start + 1]
    body = lines[start + 1 : end]
    suffix = lines[end:]

    paragraphs = split_paragraphs(body)
    abstract_blocks: list[str] = []
    keyword_lines: list[str] = []

    for paragraph in paragraphs:
        if KEYWORD_RE.match(paragraph):
            keyword_lines.append(
                "**关键词："
                + KEYWORD_RE.sub("", paragraph).replace("**", "").strip()
                + "**"
            )
            continue
        abstract_blocks.append(paragraph)

    rebuilt: list[str] = [*prefix, ""]
    for block in abstract_blocks:
        if block.strip():
            rebuilt.extend(block.splitlines())
            rebuilt.append("")
    for line in keyword_lines:
        rebuilt.append(line)
    if keyword_lines:
        rebuilt.append("")
    rebuilt.extend(suffix)
    return "\n".join(rebuilt)


def split_paragraphs(lines: list[str]) -> list[str]:
    """Combine consecutive non-empty lines into paragraphs."""
    paragraphs: list[str] = []
    buffer: list[str] = []
    for line in lines:
        if not line.strip():
            if buffer:
                paragraphs.append("\n".join(buffer).strip())
                buffer = []
            continue
        buffer.append(line.rstrip())
    if buffer:
        paragraphs.append("\n".join(buffer).strip())
    return paragraphs


def merge_image_blocks(markdown: str, work_dir: Path) -> str:
    """把每张图片规范为独立段落：前后各空一行。

    pandoc 的 implicit_figures 只对独占一段的图片生成 figure 环境
    （居中、统一图幅、带题注）；图片与文字同段会被当成行内元素，
    导致图幅不一、题注丢失、排版错乱。这里在 Markdown 层面兜底，
    写作手漏写空行时也能得到规范 figure。work_dir 仅保留调用签名。
    """
    del work_dir
    out: list[str] = []
    for line in markdown.splitlines():
        if IMAGE_BLOCK_RE.match(line):
            if out and out[-1] != "":
                out.append("")
            out.append(line.strip())
            out.append("")
            continue
        if not line.strip():
            if out and out[-1] != "":
                out.append("")
            continue
        out.append(line)
    while out and out[-1] == "":
        out.pop()
    return "\n".join(out) + "\n"


def append_source_code_appendix(markdown: str, work_dir: Path) -> str:
    """Append a compact source-code appendix."""
    if re.search(r"^\s*#\s*附录", markdown, flags=re.M) and "源代码" in markdown:
        return markdown

    candidates = collect_code_candidates(work_dir)
    if not candidates:
        return markdown

    appendix_lines = ["", "# 附录：源代码"]
    appendix_lines.append("本附录列出核心实现文件，便于复核建模流程与参数设定。")

    for idx, path in enumerate(candidates[:5], start=1):
        appendix_lines.append(f"## A.{idx} {path.name}")
        appendix_lines.append(f"文件作用：{describe_code_file(path.name)}")
        appendix_lines.append("")
        appendix_lines.append(f"```{language_for(path)}")
        appendix_lines.append(read_code_excerpt(path))
        appendix_lines.append("```")
        appendix_lines.append("")

    return markdown.rstrip() + "\n" + "\n".join(appendix_lines).rstrip() + "\n"


def collect_code_candidates(work_dir: Path) -> list[Path]:
    """Pick the most relevant code files for the appendix."""
    score_tokens = (
        "run",
        "fit",
        "predict",
        "policy",
        "search",
        "bootstrap",
        "stable",
        "core",
    )
    candidates: list[tuple[int, Path]] = []
    for path in work_dir.iterdir():
        if not path.is_file() or path.suffix.lower() not in {".m", ".py"}:
            continue
        name = path.name.lower()
        score = 0
        if path.name == "build_polished_paper.py":
            score -= 100
        if name.startswith("q2"):
            score += 90
        if (
            name.startswith("ques1")
            or name.startswith("ques2")
            or name.startswith("ques3")
            or name.startswith("ques4")
        ):
            score += 80
        score += sum(15 for token in score_tokens if token in name)
        score += min(path.stat().st_size // 20000, 20)
        candidates.append((score, path))

    candidates.sort(key=lambda item: (-item[0], item[1].name))
    return [path for _, path in candidates]


def describe_code_file(filename: str) -> str:
    """Give a short human-readable description for a code file."""
    name = filename.lower()
    if "fit" in name:
        return "模型拟合与参数估计"
    if "predict" in name:
        return "结果预测与推断"
    if "bootstrap" in name:
        return "Bootstrap 稳健性评估"
    if "policy" in name or "search" in name:
        return "策略搜索与决策选择"
    if "stable" in name or "core" in name:
        return "核心求解流程"
    return "核心实现文件"


def language_for(path: Path) -> str:
    """Map file suffix to code fence language."""
    if path.suffix.lower() == ".m":
        return "matlab"
    return "python"


def read_code_excerpt(path: Path, max_chars: int = 4000) -> str:
    """Read a bounded excerpt of a code file for the appendix."""
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rstrip() + "\n..."


def ensure_blank_lines_around_headings(markdown: str) -> str:
    """Keep blank lines around headings for stable Markdown parsing."""
    lines = markdown.splitlines()
    normalized: list[str] = []
    for idx, line in enumerate(lines):
        stripped = line.rstrip()
        if HEADING_RE.match(stripped):
            if normalized and normalized[-1].strip():
                normalized.append("")
            normalized.append(stripped)
            next_line = lines[idx + 1].strip() if idx + 1 < len(lines) else ""
            if next_line and not HEADING_RE.match(next_line):
                normalized.append("")
        else:
            normalized.append(line)
    return "\n".join(normalized)


def build_reference_docx(work_dir: Path) -> Path:
    """Create a reference DOCX used by pandoc as a style template."""
    reference_path = work_dir / "paper_reference.docx"
    doc = Document()
    configure_document(doc)
    doc.save(reference_path)
    return reference_path


def configure_document(doc: Document) -> None:
    """Configure page size, margins, fonts and footer."""
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.2)

    styles = doc.styles
    set_style_font(styles["Normal"], east="SimSun", latin="Times New Roman", size=12)
    normal = styles["Normal"].paragraph_format
    normal.first_line_indent = Pt(24)
    normal.line_spacing = 1.32
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    style_specs = [
        ("Title", "SimHei", 17.5, 42, 18, WD_ALIGN_PARAGRAPH.CENTER, True),
        ("Heading 1", "SimHei", 15, 12, 8, WD_ALIGN_PARAGRAPH.CENTER, True),
        ("Heading 2", "SimHei", 13.5, 8, 5, WD_ALIGN_PARAGRAPH.LEFT, True),
        ("Heading 3", "SimHei", 12.5, 6, 3, WD_ALIGN_PARAGRAPH.LEFT, True),
        ("Caption", "SimSun", 10.5, 6, 4, WD_ALIGN_PARAGRAPH.CENTER, False),
    ]
    for name, east, size, before, after, align, bold in style_specs:
        style = styles[name]
        set_style_font(style, east=east, latin="Times New Roman", size=size, bold=bold)
        p = style.paragraph_format
        p.first_line_indent = Pt(0)
        p.line_spacing = 1.25
        p.space_before = Pt(before)
        p.space_after = Pt(after)
        p.alignment = align
        p.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    add_page_field(footer)
    for run in footer.runs:
        set_run_font(run, east="SimSun", latin="Times New Roman", size=10.5)


def add_page_field(paragraph) -> None:
    """Insert a PAGE field in the footer."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def set_style_font(style, east="SimSun", latin="Times New Roman", size=12, bold=False):
    """Set font for a Word style."""
    font = style.font
    font.name = latin
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string("000000")
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def set_run_font(run, east="SimSun", latin="Times New Roman", size=12, bold=None):
    """Set font for a Word run."""
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def convert_markdown_to_docx(
    markdown: str,
    output_path: Path,
    reference_docx: Path,
    resource_path: Path,
) -> None:
    """Convert Markdown to DOCX using pandoc."""
    convert_text(
        markdown,
        to="docx",
        format="markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_html",
        outputfile=str(output_path),
        extra_args=[
            f"--reference-doc={reference_docx}",
            f"--resource-path={resource_path}",
            "--standalone",
            "--wrap=none",
        ],
    )


def convert_markdown_to_pdf(
    markdown: str,
    output_path: Path,
    resource_path: Path,
) -> None:
    """Convert Markdown to PDF as a best-effort deliverable."""
    header_path = resource_path / "paper_pdf_header.tex"
    header_path.write_text(build_pdf_header(resource_path), encoding="utf-8")

    try:
        convert_text(
            markdown,
            to="pdf",
            format="markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_html",
            outputfile=str(output_path),
            extra_args=[
                f"--resource-path={resource_path}",
                f"--include-in-header={header_path}",
                "--standalone",
                "--wrap=none",
                "--pdf-engine=xelatex",
            ],
        )
        return
    except WorkCancelled:
        raise
    except Exception as exc:
        logger.warning(
            "pandoc pdf export failed, falling back to HTML/WeasyPrint: %s", exc
        )

    html = convert_text(
        markdown,
        to="html5",
        format="markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_html",
        extra_args=[
            f"--resource-path={resource_path}",
            "--standalone",
            "--wrap=none",
            "--mathml",
        ],
    )

    from weasyprint import CSS, HTML  # type: ignore[import-unresolved]

    html_path = resource_path / "paper_render.html"
    html_path.write_text(html, encoding="utf-8")
    HTML(string=html, base_url=str(resource_path)).write_pdf(
        str(output_path),
        stylesheets=[CSS(string=build_weasyprint_css())],
    )


def build_weasyprint_css() -> str:
    """Build a lightweight CSS stylesheet for HTML-to-PDF fallback."""
    return """
@page {
  size: A4;
  margin: 2.6cm 2.7cm 2.2cm 2.7cm;
  @bottom-center {
    content: counter(page);
    font-family: "Times New Roman", "SimSun";
    font-size: 10.5pt;
  }
}

body {
  font-family: "Times New Roman", "SimSun";
  font-size: 12pt;
  line-height: 1.32;
  text-align: justify;
  color: #000;
}

h1, h2, h3, h4, h5, h6 {
  page-break-after: avoid;
  page-break-inside: avoid;
}

h1 {
  font-size: 17.5pt;
  text-align: center;
  font-weight: 700;
}

h2 {
  font-size: 15pt;
  text-align: center;
  font-weight: 700;
}

h3 {
  font-size: 13.5pt;
  font-weight: 700;
}

p {
  margin: 0 0 0.45em 0;
}

img {
  max-width: 100%;
  height: auto;
}

table {
  border-collapse: collapse;
  width: 100%;
}

th, td {
  border: 1px solid #444;
  padding: 0.22em 0.35em;
}

code, pre {
  white-space: pre-wrap;
  word-break: break-word;
}
"""


def build_pdf_header(
    resource_path: Path,
    comp_template: CompTemplate = CompTemplate.CHINA,
) -> str:
    """Create a portable XeLaTeX header for the selected competition."""
    paper = "a4paper" if comp_template == CompTemplate.CHINA else "letterpaper"
    if (resource_path / "simhei.ttf").is_file():
        cjk_font = r"""\setCJKmainfont[
  Path=./,
  BoldFont=simhei.ttf
]{simhei.ttf}"""
    else:
        # Docker 使用系统字体；Fandol 兼容未安装 Noto 的 TeX 发行版。
        cjk_font = r"""\IfFontExistsTF{Noto Sans CJK SC}{
  \setCJKmainfont{Noto Sans CJK SC}
}{
  \setCJKmainfont[BoldFont=FandolSong-Bold]{FandolSong-Regular}
}"""
    return f"""
\\usepackage[{paper},margin=2.54cm]{{geometry}}
\\usepackage{{fontspec}}
\\usepackage{{xeCJK}}
\\usepackage{{amsmath,amssymb,bm}}
\\usepackage{{booktabs,longtable,array,graphicx,float}}
\\IfFontExistsTF{{Times New Roman}}{{\\setmainfont{{Times New Roman}}}}{{\\setmainfont{{TeX Gyre Termes}}}}
{cjk_font}
\\IfFontExistsTF{{Arial}}{{\\setsansfont{{Arial}}}}{{\\setsansfont{{TeX Gyre Heros}}}}
\\IfFontExistsTF{{Consolas}}{{\\setmonofont{{Consolas}}}}{{\\setmonofont{{Latin Modern Mono}}}}
\\XeTeXlinebreaklocale "zh"
\\XeTeXlinebreakskip = 0pt plus 1pt
\\setlength{{\\emergencystretch}}{{3em}}
\\setlength{{\\parindent}}{{2em}}
\\setlength{{\\parskip}}{{0.25em}}
"""
