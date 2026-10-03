import io
import zipfile

import pytest

from app.utils.document_parser import save_word_preview_images
from app.utils.pdf_parser import PdfParseError


def test_word_import_keeps_image_and_rewrites_preview_url(tmp_path, monkeypatch):
    import hashlib

    import fitz
    from docx import Document

    from app.config import setting
    from app.utils.document_parser import parse_word_bytes

    with fitz.open() as pdf:
        page = pdf.new_page(width=10, height=10)
        png = page.get_pixmap().tobytes("png")
    document = Document()
    document.add_paragraph("含图片的赛题")
    document.add_picture(io.BytesIO(png))
    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()
    monkeypatch.setattr(setting, "BACKEND_ROOT", tmp_path)
    parsed = parse_word_bytes(content, ".docx")
    digest = hashlib.sha256(content).hexdigest()
    assert f"](/static/_document_previews/{digest}/image1.png)" in parsed.text
    assert "含图片的赛题" in parsed.text
    assert (
        tmp_path / "project" / "work_dir" / "_document_previews" / digest / "image1.png"
    ).read_bytes() == png


def archive(entries):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as output:
        for name, content in entries.items():
            output.writestr(name, content)
    return stream.getvalue()


def test_rasterizes_svg_and_skips_paths_outside_media(tmp_path):
    data = archive(
        {
            "word/media/image2.svg": '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10" fill="red"/></svg>',
            "word/media/../../escape.png": b"invalid",
        }
    )
    mapping = save_word_preview_images(data, tmp_path / "media")
    assert mapping == {"media/image2.svg": "image2.svg.png"}
    image = tmp_path / "media/image2.svg.png"
    assert image.read_bytes().startswith(b"\x89PNG")
    assert not (tmp_path / "escape.png").exists()
    assert save_word_preview_images(data, tmp_path / "media") == mapping


def test_does_not_overwrite_existing_file(tmp_path):
    (tmp_path / "image.png").write_bytes(b"original")
    with pytest.raises(PdfParseError, match="冲突"):
        save_word_preview_images(archive({"word/media/image.png": b"other"}), tmp_path)
    assert (tmp_path / "image.png").read_bytes() == b"original"
