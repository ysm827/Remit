"""独立论文工作区的数据安全、编译快照和流程解耦回归。"""

import asyncio
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.core.attachment_profile import profile_attachment
from app.core.data_scout import build_data_profile
from app.core.workflow_checkpoint import WorkflowCheckpoint
from app.routers import writing_router
from app.schemas.request import Problem
from app.services import writing_workspace as workspace
from app.utils.document_parser import parse_word_bytes
from app.utils.pdf_parser import PdfParseError


def test_word_preserves_paragraph_table_order():
    document = Document()
    document.add_paragraph("赛题背景")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "温度"
    row.cells[1].text = "42"
    document.add_paragraph("问题一：分析变化")
    data = BytesIO()
    document.save(data)
    parsed = parse_word_bytes(data.getvalue(), ".docx")
    assert (
        parsed.text.index("背景")
        < parsed.text.index("温度 | 42")
        < parsed.text.index("问题一")
    )
    assert parsed.char_count == len(parsed.text)
    assert parsed.page_count == 0


def test_word_rejects_fake_document():
    with pytest.raises(PdfParseError):
        parse_word_bytes(b"fake", ".docx")


def test_attachment_profile_and_unknown_inputs(tmp_path):
    (tmp_path / "data.json").write_text('[{"x": 3}]', encoding="utf-8")
    (tmp_path / "sample.custom").write_bytes(b"custom format")
    (tmp_path / ".remit-inputs.json").write_text(
        '["data.json", "sample.custom"]', encoding="utf-8"
    )
    profile = build_data_profile(tmp_path)
    assert set(profile["discovered_files"]) == {"data.json", "sample.custom"}
    by_name = {item["file"]: item for item in profile["files"]}
    assert by_name["data.json"]["status"] == "parsed"
    assert by_name["sample.custom"]["status"] == "metadata_only"


def test_numpy_does_not_unpickle_objects(tmp_path):
    import numpy as np

    path = tmp_path / "object.npy"
    np.save(path, {"untrusted": True})
    with pytest.raises(ValueError):
        profile_attachment(path)


def test_new_workflow_has_no_writer_nodes(tmp_path):
    checkpoint = WorkflowCheckpoint(tmp_path)
    state = checkpoint.initialize(Problem(task_id="separate"))
    state.update(questions={"ques1": "预测"}, modeler_response={})
    order = checkpoint.node_order(state)
    assert order[-1] == "sync_writing"
    assert not any(node.startswith("write:") or node == "finalize" for node in order)
    state["workflow_features"].remove("separate_writing")
    assert checkpoint.node_order(state)[-1] == "finalize"


def test_sync_preserves_user_source_and_copies_evidence(tmp_path):
    (tmp_path / "plot.png").write_bytes(b"image")
    state = {
        "solution_results": {
            "ques1": {"artifacts": ["plot.png"], "paper_ready_images": []}
        }
    }
    root = workspace.ensure_workspace(tmp_path)
    (root / "main.tex").write_text("manual edit", encoding="utf-8")
    first = workspace.sync_results(tmp_path, state)
    (tmp_path / "plot.png").unlink()
    assert (
        root / ".inputs" / first["revision"] / "assets" / "plot.png"
    ).read_bytes() == b"image"
    assert (root / "main.tex").read_text(encoding="utf-8") == "manual edit"


def test_save_conflict_history_and_path_traversal(tmp_path):
    root = workspace.ensure_workspace(tmp_path)
    original = (root / "main.tex").read_text(encoding="utf-8")
    new_version = workspace.save_source(
        root, "main.tex", "changed", workspace.digest(original)
    )
    assert new_version == workspace.digest("changed")
    assert len(list((root / ".history").glob("*/source.txt"))) == 1
    with pytest.raises(FileExistsError):
        workspace.save_source(root, "main.tex", "overwrite", workspace.digest(original))
    for name in ("../outside.tex", "C:/outside.tex", ".inputs/input.tex", "/root.tex"):
        with pytest.raises(ValueError):
            workspace.resolve_source(root, name)


def test_compile_snapshot_is_immutable(tmp_path):
    root = workspace.ensure_workspace(tmp_path)
    build, main, revision = workspace.prepare_build(root)
    (root / "main.tex").write_text("changed during compile", encoding="utf-8")
    assert (build / main).read_text(encoding="utf-8") == workspace.TEMPLATE
    assert workspace.project_revision(root) != revision


@pytest.fixture
def writing_client(tmp_path, monkeypatch):
    checkpoint = WorkflowCheckpoint(tmp_path)
    state = checkpoint.initialize(Problem(task_id="paper-test"))
    state["status"] = "completed"
    checkpoint.save(state)
    monkeypatch.setattr(
        writing_router, "_resolve_task_directory", lambda task_id: tmp_path
    )
    writing_router._locks.clear()
    writing_router._compiling.clear()
    app = FastAPI()
    app.include_router(writing_router.router)
    with TestClient(app) as client:
        yield client, tmp_path


def test_failed_compile_retains_last_pdf(writing_client, monkeypatch):
    client, task_root = writing_client
    client.get("/api/writing/paper-test")
    root = workspace.paper_root(task_root)
    (root / "preview.pdf").write_bytes(b"last valid PDF")
    workspace.write_json(root / "compile.json", {"pdf_revision": "previous"})
    monkeypatch.setattr(
        workspace, "compile_build", lambda *args: {"status": "failed", "log": "error"}
    )
    response = client.post("/api/writing/paper-test/compile")
    assert response.status_code == 200
    assert response.json()["pdf_revision"] == "previous"
    assert (root / "preview.pdf").read_bytes() == b"last valid PDF"


def test_writing_routes_block_unfinished_results_and_version_conflicts(writing_client):
    client, task_root = writing_client
    client.get("/api/writing/paper-test")
    source = client.get("/api/writing/paper-test/source").json()
    assert (
        client.post(
            "/api/writing/paper-test/source",
            json={"name": "main.tex", "content": "new", "version": source["version"]},
        ).status_code
        == 200
    )
    assert (
        client.post(
            "/api/writing/paper-test/source",
            json={"name": "main.tex", "content": "old", "version": source["version"]},
        ).status_code
        == 409
    )
    checkpoint = WorkflowCheckpoint(task_root)
    checkpoint.mark_status("stopped")
    assert client.post("/api/writing/paper-test/sync").status_code == 409
    assert client.post("/api/writing/paper-test/generate").status_code == 409


def test_writer_failure_does_not_modify_modeling_checkpoint(writing_client):
    from app.schemas.response import SystemMessage
    from app.services.redis_manager import redis_manager

    _, task_root = writing_client
    root = workspace.ensure_workspace(task_root)
    state = WorkflowCheckpoint(task_root).load()
    state["questions"] = {"background": "background", "ques1": "question"}
    state["ques_count"] = 1
    state["solution_results"] = {"ques1": {"writer_prompt": "write", "artifacts": []}}
    inputs = workspace.sync_results(task_root, state)
    before = (task_root / "workflow_state.json").read_bytes()

    async def failing_writer(*args, **kwargs):
        await redis_manager.publish_message(
            "paper-test", SystemMessage(content="writer failure", type="error")
        )
        raise RuntimeError("model offline")

    with (
        patch(
            "app.core.llm.llm_factory.LLMFactory.get_writer_llm", return_value=object()
        ),
        patch(
            "app.core.agents.writer_agent.WriterAgent.run",
            new=AsyncMock(side_effect=failing_writer),
        ),
    ):
        asyncio.run(writing_router._generate("paper-test", root, inputs["revision"]))
    assert (
        workspace.read_json(root / "workspace.json")["generation"]["status"] == "failed"
    )
    assert (task_root / "workflow_state.json").read_bytes() == before
    assert list((root / ".drafts").glob("*/events.jsonl"))
    assert not redis_manager.archive.exists("paper-test")


@pytest.mark.parametrize("edit_during_generation", [False, True])
def test_writer_resume_reuses_validated_sections(writing_client, edit_during_generation):
    from app.schemas.A2A import WriterResponse

    client, task_root = writing_client
    client.get("/api/writing/paper-test")
    root = workspace.paper_root(task_root)
    state = WorkflowCheckpoint(task_root).load()
    state["ques_count"] = 1
    state["questions"] = {"ques1": "question"}
    state["solution_results"] = {
        "eda": {"writer_prompt": "eda", "artifacts": []},
        "ques1": {"writer_prompt": "question", "artifacts": []},
    }
    inputs = workspace.sync_results(task_root, state)
    writer = AsyncMock(
        side_effect=[
            WriterResponse(response_content="保存过的真实章节"),
            RuntimeError("technical interruption"),
            WriterResponse(response_content="恢复后完成的第二章"),
        ]
    )

    def convert(markdown, path, *_args):
        path.write_text(markdown, encoding="utf-8")
        if edit_during_generation:
            (root / "main.tex").write_text("new manual edits", encoding="utf-8")

    with (
        patch.object(writing_router, "compile_source", AsyncMock(return_value={"status": "completed"})) as compiler,
        patch(
            "app.models.user_output.UserOutput._section_order",
            return_value=["eda", "ques1"],
        ),
        patch(
            "app.core.llm.llm_factory.LLMFactory.get_writer_llm", return_value=object()
        ),
        patch("app.core.agents.writer_agent.WriterAgent.run", writer),
        patch("app.core.deliverable_contract.validate_writer_section"),
        patch("app.core.flows.Flows.get_write_flows", return_value={}),
        patch(
            "app.utils.paper_polish.polish_markdown", side_effect=lambda text, *_: text
        ),
        patch("app.utils.paper_polish._convert_markdown_to_latex", side_effect=convert),
        patch(
            "app.services.competitions.adapt_generated_source",
            side_effect=lambda _, text: text,
        ),
    ):
        asyncio.run(writing_router._generate("paper-test", root, inputs["revision"]))
        first = workspace.read_json(root / "workspace.json")["generation"]
        assert first["status"] == "failed"
        assert first["partial"] is True
        assert first["completed_sections"] == ["eda"]
        assert (root / first["file"]).is_file()
        asyncio.run(writing_router._generate("paper-test", root, inputs["revision"]))
    second = workspace.read_json(root / "workspace.json")["generation"]
    assert second["status"] == "completed"
    assert second["generation_id"] == first["generation_id"]
    assert writer.await_count == 3
    if edit_during_generation:
        assert workspace.read_json(root / "workspace.json")["main"] == "main.tex"
        assert (root / "main.tex").read_text(encoding="utf-8") == "new manual edits"
        compiler.assert_not_awaited()
    else:
        assert workspace.read_json(root / "workspace.json")["main"] == second["file"]
        assert compiler.await_count == 3  # first chapter, resumed second chapter, final
    assert "保存过的真实章节" in (root / second["file"]).read_text(encoding="utf-8")


@pytest.mark.parametrize("repair_succeeds", [True, False])
def test_writer_repairs_rejected_section_once_without_publishing_invalid_prose(writing_client, repair_succeeds):
    from app.schemas.A2A import WriterResponse

    client, task_root = writing_client
    client.get("/api/writing/paper-test")
    root = workspace.paper_root(task_root)
    state = WorkflowCheckpoint(task_root).load()
    state["solution_results"] = {"eda": {
        "writer_prompt": "撰写数据分析章节", "artifacts": [],
        "quality_report": {"selected_model": "指定方法"},
    }}
    inputs = workspace.sync_results(task_root, state)
    rejected = "已读取数据并核对字段和单位。" * 30
    repaired = rejected + ("采用指定方法。" if repair_succeeds else "")
    writer = AsyncMock(side_effect=[WriterResponse(response_content=rejected), WriterResponse(response_content=repaired)])
    with (
        patch("app.core.llm.llm_factory.LLMFactory.get_writer_llm", return_value=object()),
        patch("app.models.user_output.UserOutput._section_order", return_value=["eda"]),
        patch("app.core.agents.writer_agent.WriterAgent.run", writer),
        patch("app.core.flows.Flows.get_write_flows", return_value={}),
        patch("app.utils.paper_polish.polish_markdown", side_effect=lambda text, *_: text),
        patch("app.utils.paper_polish._convert_markdown_to_latex", side_effect=lambda text, path, *_: path.write_text(text, encoding="utf-8")),
        patch("app.services.competitions.adapt_generated_source", side_effect=lambda _, text: text),
        patch.object(writing_router, "compile_source", AsyncMock(return_value={"status": "completed"})),
    ):
        asyncio.run(writing_router._generate("paper-test", root, inputs["revision"]))
    generation = workspace.read_json(root / "workspace.json")["generation"]
    assert writer.await_count == 2
    assert "未说明质量报告中的入选模型" in writer.call_args_list[1].args[0]
    assert generation["status"] == ("completed" if repair_succeeds else "failed")
    assert bool(list(root.glob("draft-*.tex"))) == repair_succeeds
    attempt = workspace.read_json(next((root / ".drafts").glob("*/.attempts/eda.json")))
    assert attempt["attempt"] == 2
    assert attempt["response"]["response_content"] == repaired

    if repair_succeeds:
        return
    # Simulate a corrected validator recognizing the saved response; no new model
    # call is needed, but the retained candidate must be validated again.
    with (
        patch("app.core.llm.llm_factory.LLMFactory.get_writer_llm", return_value=object()),
        patch("app.models.user_output.UserOutput._section_order", return_value=["eda"]),
        patch("app.core.agents.writer_agent.WriterAgent.run", AsyncMock()) as resumed,
        patch("app.core.deliverable_contract.validate_writer_section") as validate,
        patch("app.core.flows.Flows.get_write_flows", return_value={}),
        patch("app.utils.paper_polish.polish_markdown", side_effect=lambda text, *_: text),
        patch("app.utils.paper_polish._convert_markdown_to_latex", side_effect=lambda text, path, *_: path.write_text(text, encoding="utf-8")),
        patch("app.services.competitions.adapt_generated_source", side_effect=lambda _, text: text),
        patch.object(writing_router, "compile_source", AsyncMock(return_value={"status": "completed"})),
    ):
        asyncio.run(writing_router._generate("paper-test", root, inputs["revision"]))
    resumed.assert_not_awaited()
    validate.assert_called_once()
    assert validate.call_args.args[:2] == ("eda", repaired)
    assert workspace.read_json(root / "workspace.json")["generation"]["status"] == "completed"


def test_generation_resume_is_limited_to_failed_same_input(tmp_path):
    root = workspace.ensure_workspace(tmp_path)
    for status, revision, identifier, reuse in [
        ("failed", "rev", "123456abcd", True),
        ("interrupted", "rev", "123456abcd", True),
        ("completed", "rev", "123456abcd", False),
        ("failed", "older", "123456abcd", False),
        ("failed", "rev", "../outside", False),
    ]:
        workspace.write_json(
            root / "workspace.json",
            {
                "generation": {
                    "status": status,
                    "input_revision": revision,
                    "generation_id": identifier,
                }
            },
        )
        chosen = writing_router._generation_id(root, "rev")
        assert (chosen == identifier) is reuse
